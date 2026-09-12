import io
import json
import shutil
import zipfile

import pytest
from fastapi.testclient import TestClient

from app import store
from app.main import app
from app.studio.extractors import extract, safe_name, zip_members
from app.studio.network import public_target, repository_url
from app.studio.pipeline import prepare
from app.studio.schemas import Recipe


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'ROOT', tmp_path)
    return TestClient(app)


def recipe(**kwargs):
    return Recipe(source_ids=['source'], min_chars=1, **kwargs)


@pytest.mark.parametrize('name,raw,needle',[
    ('readme.md', b'# Useful documentation', 'Useful'),
    ('page.html', b'<nav>MENU</nav><main>Hello world</main><script>BAD</script>', 'Hello'),
    ('data.xml', b'<root><doc>Example</doc></root>', 'Example'),
    ('data.yaml', b'name: example\ntext: content', 'example'),
    ('main.py', b'def add(a,b):\n    return a+b', '    return'),
    ('data.json', b'[{"text":"one"},{"text":"two"}]', 'one'),
    ('data.jsonl', b'{"text":"one"}\n{"text":"two"}', 'one'),
    ('data.csv', b'text\none\ntwo', 'one'),
])
def test_formats(name, raw, needle):
    rows, _, _ = extract(raw, name, recipe())
    assert needle in rows[0]['text']
    if name.endswith('.html'):
        assert 'MENU' not in rows[0]['text'] and 'BAD' not in rows[0]['text']


def test_docx_tables_and_pdf():
    from docx import Document
    from pypdf import PdfWriter
    doc=Document(); doc.add_paragraph('Paragraph'); table=doc.add_table(rows=1, cols=1);table.cell(0,0).text='Cell'
    stream=io.BytesIO();doc.save(stream)
    assert 'Cell' in extract(stream.getvalue(),'test.docx',recipe())[0][0]['text']
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);stream=io.BytesIO();writer.write(stream)
    with pytest.raises(ValueError,match='OCR'):
        extract(stream.getvalue(),'empty.pdf',recipe())


@pytest.mark.parametrize('name',['../escape.py','/etc/passwd','C:\\secret.txt','x/../../bad','bad\x00.txt'])
def test_paths(name):
    with pytest.raises(ValueError):safe_name(name)


def test_zip_limits_and_ignored():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as archive:
        archive.writestr('repo/main.py','print(1)')
        archive.writestr('repo/node_modules/x.js','ignored')
    entries=list(zip_members(stream.getvalue(),recipe()))
    assert entries[0][1]==b'print(1)' and entries[1][1] is None
    with pytest.raises(ValueError):list(zip_members(stream.getvalue(),recipe(max_files=1)))
    with pytest.raises(zipfile.BadZipFile):list(zip_members(b'not zip',recipe()))
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as archive:archive.writestr('../bad.txt','bad')
    with pytest.raises(ValueError):list(zip_members(stream.getvalue(),recipe()))


def test_xml_entities_and_binary():
    with pytest.raises(Exception):extract(b'<!DOCTYPE x [<!ENTITY y SYSTEM "file:///etc/passwd">]><x>&y;</x>','a.xml',recipe())
    with pytest.raises(ValueError):extract(b'MZ\x00payload','fake.py',recipe())


@pytest.mark.parametrize('url',['http://example.com','https://127.0.0.1','https://[::1]','file:///etc/passwd','https://user:password@example.com','https://example.com?token=secret'])
def test_network_rejects_private_and_credentials(url):
    with pytest.raises(ValueError):public_target(url)


def test_dns_all_addresses_validated(monkeypatch):
    monkeypatch.setattr('socket.getaddrinfo',lambda *a,**k:[(2,1,6,'',('93.184.216.34',443)),(2,1,6,'',('10.1.1.1',443))])
    with pytest.raises(ValueError):public_target('https://example.com')
    assert repository_url('https://github.com/owner/repo','main')=='https://codeload.github.com/owner/repo/zip/main'


def test_prepare_publish_split_integrity_and_resume(client):
    sources=[]
    for i in range(8):
        text=f'Documento {i} sobre tópico {i}. '+(f'Conteúdo autoral {i} para estudar modelos {i} e preparar dados {i}. '*5)
        result=client.post('/api/studio/sources',files={'files':(f'file{i}.txt',text.encode(),'text/plain')})
        assert result.status_code==200,result.text
        sources.append(result.json()[0]['id'])
    response=client.post('/api/studio/prepare',json={'name':'Teste Data Studio','source_ids':sources,'chunk_chars':128,'min_words':1})
    assert response.status_code==200,response.text
    job=response.json();folder=store.run_dir(job['id'])
    prepare(job,folder)
    snapshot=(folder/'preview.json').read_bytes()
    prepare(job,folder)
    assert (folder/'preview.json').read_bytes()==snapshot
    store.update('jobs',job['id'],{'status':'completed'})
    preview=client.get('/api/studio/preparations/'+job['id']).json()
    assert preview['manifest']['counts']['test']>0
    response=client.post('/api/studio/preparations/'+job['id']+'/publish',json={})
    assert response.status_code==200,response.text
    dataset=response.json()
    assert client.post('/api/studio/preparations/'+job['id']+'/publish',json={}).json()['id']==dataset['id']
    assert client.get('/api/studio/datasets/'+dataset['id']+'/dataset.jsonl').status_code==200
    assert client.post('/api/datasets/import',data={'name':'edit','dataset_id':dataset['id'],'content':'New text'}).status_code==422
    train=client.post('/api/train',json={'name':'smoke','mode':'scratch','dataset_id':dataset['id'],'max_steps':1,'context':128})
    assert train.status_code==200,train.text
    training=json.loads((store.run_dir(train.json()['id'])/'dataset.json').read_text())
    assert not {r['group'] for r in training['train']} & {r['group'] for r in training['validation']}
    assert all(r['split']!='test' for values in training.values() for r in values)
    assert client.post('/api/train',json={'name':'bad','mode':'qlora','model_id':'missing','dataset_id':dataset['id']}).status_code==422


def test_conversation_fields_and_no_synthetic_answers():
    rows,_,_=extract(b'[{"ask":"Question","answer":"Answer"}]','chat.json',recipe(preset='conversations',instruction_field='ask',output_field='answer'))
    assert rows[0]['messages'][-1]['content']=='Answer'
    with pytest.raises(ValueError):extract(b'[{"random":"document"}]','chat.json',recipe(preset='conversations'))


def test_empty_data_and_validation(client):
    assert client.post('/api/studio/sources',data={}).status_code==422
    assert client.post('/api/studio/prepare',json={'source_ids':[]}).status_code==422
    assert client.post('/api/studio/prepare',json={'source_ids':['x'],'chunk_chars':128,'overlap':128}).status_code==422
    assert client.post('/api/studio/prepare',json={'source_ids':['x'],'overlap_words':180,'max_words':180}).status_code==422


@pytest.mark.parametrize('dedupe', [True, False])
def test_chunk_duplicates_never_cross_splits(client, dedupe):
    from app.studio.storage import save_source
    sources=[]
    shared='Shared paragraph about maintaining equipment and keeping a reliable record of maintenance tasks.'
    for text in [shared+'\n\n'+'Unique part A about mathematics and proofs. '*3, shared+'\n\n'+'Unique part B about geology and mountains. '*3, 'Independent third document about stars.', 'Independent fourth document about cooking.']:
        sources.append(save_source(f'file{len(sources)}.txt',text.encode())['id'])
    job=client.post('/api/studio/prepare',json={'source_ids':sources,'min_chars':1,'min_words':1,'quality_filter':'off','chunking':'paragraph','chunk_chars':128,'exact_dedupe':dedupe}).json()
    prepare(job,store.run_dir(job['id']))
    rows=json.loads((store.run_dir(job['id'])/'preview.json').read_text())
    assert len({r['split'] for r in rows})==3
    related=[r for r in rows if r['metadata']['path'] in ('file0.txt','file1.txt')]
    assert len({r['split'] for r in related})==1
    if dedupe:assert len({r['text'] for r in rows})==len(rows)


def test_near_dedupe_and_resume_api(client):
    from app.studio.storage import save_source
    texts=['alpha beta gamma delta epsilon zeta eta theta iota kappa',
           'alpha beta gamma delta epsilon zeta eta theta iota kappa extra',
           'Space telescopes capture distant stars and galaxies beyond our solar system.',
           'A bicycle uses wheels and pedals and a chain to travel along roads.']
    sources=[save_source(f'file{i}.txt',text.encode())['id'] for i,text in enumerate(texts)]
    job=client.post('/api/studio/prepare',json={'source_ids':sources,'min_chars':1,'min_words':1,'quality_filter':'off','near_dedupe':True,'similarity':0.8}).json()
    folder=store.run_dir(job['id']);prepare(job,folder)
    manifest=json.loads((folder/'data-manifest.json').read_text())
    assert manifest['duplicates']==1
    store.update('jobs',job['id'],{'status':'interrupted'})
    assert client.post('/api/studio/preparations/'+job['id']+'/resume').json()['status']=='queued'
    assert client.post('/api/studio/preparations/'+job['id']+'/publish',json={}).status_code==422


def test_streamed_body_limit_without_content_length():
    import asyncio
    from app.studio.limits import BodyLimit
    from starlette.exceptions import HTTPException
    async def receiver():return {'type':'http.request','body':b'x'*1024**2,'more_body':True}
    async def consumer(scope,receive,send):
        for _ in range(35):await receive()
    async def sender(message):pass
    with pytest.raises(HTTPException) as error:
        asyncio.run(BodyLimit(consumer)({'type':'http','path':'/api/datasets/import'},receiver,sender))
    assert error.value.status_code==413


def test_versioning_preserves_original(client):
    from app.studio.storage import save_source
    sources=[save_source(f'v{i}.txt',f'Independent document {i} with distinct facts number {i}.'.encode())['id'] for i in range(4)]
    first=None
    for version in (1,2):
        config={'source_ids':sources,'name':'Versioned','parent_id':first['id'] if first else None,'min_words':1,'min_chars':1,'quality_filter':'off'}
        job=client.post('/api/studio/prepare',json=config).json()
        folder=store.run_dir(job['id']);prepare(job,folder);store.update('jobs',job['id'],{'status':'completed'})
        dataset=client.post('/api/studio/preparations/'+job['id']+'/publish',json={}).json()
        assert dataset['version']==version
        if first is None:first=dataset
        else:
            assert dataset['id']!=first['id']
            assert client.get('/api/datasets/'+first['id']).json()['sha256']==first['sha256']


@pytest.mark.skipif(not shutil.which('tesseract') or not shutil.which('pdftoppm'), reason='OCR tools are available in the Docker image')
def test_real_scanned_pdf_ocr():
    Image = pytest.importorskip('PIL.Image')
    from PIL import ImageDraw, ImageFont
    image=Image.new('RGB',(1400,300),'white')
    draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40)
    draw.text((40,80),'Training data needs careful review.',fill='black',font=font)
    stream=io.BytesIO();image.save(stream,format='PDF',resolution=150)
    rows,kind,_=extract(stream.getvalue(),'scan.pdf',recipe(ocr=True,ocr_language='eng'))
    assert kind=='pdf' and 'careful review' in rows[0]['text']


def test_pdf_pages_are_stitched(monkeypatch):
    pypdf = pytest.importorskip('pypdf')
    class Page:
        def __init__(self, text):
            self._text = text
        def extract_text(self):
            return self._text
    class Reader:
        def __init__(self, *_a, **_k):
            self.pages = [Page('Hello from page one that continues'), Page('on the second page with more words for training.')]
    monkeypatch.setattr(pypdf, 'PdfReader', Reader)
    rows, kind, _ = extract(b'%PDF-fake', 'book.pdf', recipe())
    assert kind == 'pdf' and len(rows) == 1
    assert 'page one' in rows[0]['text'] and 'second page' in rows[0]['text']
    assert rows[0]['pages'] == [1, 2]


def test_long_pdf_is_extracted(monkeypatch):
    pypdf = pytest.importorskip('pypdf')
    class Page:
        def extract_text(self):
            return 'Uma frase literária com palavras suficientes para treinar o modelo de linguagem.'
    class Reader:
        def __init__(self, *_a, **_k):
            self.pages = [Page() for _ in range(600)]
    monkeypatch.setattr(pypdf, 'PdfReader', Reader)
    rows, kind, _ = extract(b'%PDF-fake', 'obras.pdf', recipe())
    assert kind == 'pdf' and len(rows[0]['pages']) == 600
    assert 'warning' not in rows[0]


def test_pdf_page_cap_truncates_with_warning(monkeypatch):
    pypdf = pytest.importorskip('pypdf')
    class Page:
        def extract_text(self):
            return 'Texto de uma página extraída do volume.'
    class Reader:
        def __init__(self, *_a, **_k):
            self.pages = [Page() for _ in range(40)]
    monkeypatch.setattr(pypdf, 'PdfReader', Reader)
    rows, _, _ = extract(b'%PDF-fake', 'obras.pdf', recipe(max_pdf_pages=12))
    assert rows[0]['pages'] == list(range(1, 13))
    assert '12' in rows[0]['warning']


def test_extraction_error_is_surfaced(client):
    from app.studio.storage import save_source
    source = save_source('book.pdf', b'not-a-pdf')
    job = client.post('/api/studio/prepare', json={'source_ids': [source['id']], 'min_chars': 1, 'min_words': 1, 'quality_filter': 'off'}).json()
    with pytest.raises(ValueError, match='Assinatura PDF'):
        prepare(job, store.run_dir(job['id']))


def test_single_book_gets_sequential_splits(client):
    from app.studio.storage import save_source
    chapters = '\n\n'.join(prose(12, i) for i in range(24))
    source = save_source('machado.txt', chapters.encode())
    job = client.post('/api/studio/prepare', json={
        'source_ids': [source['id']], 'min_chars': 1, 'min_words': 8, 'quality_filter': 'off',
        'chunking': 'sentence', 'max_words': 40
    }).json()
    prepare(job, store.run_dir(job['id']))
    folder = store.run_dir(job['id'])
    rows = json.loads((folder/'preview.json').read_text())
    assert {row['split'] for row in rows} == {'train', 'validation', 'test'}
    ordered = sorted(rows, key=lambda r: r['metadata']['chunk'])
    rank = {'train': 0, 'validation': 1, 'test': 2}
    labels = [row['split'] for row in ordered]
    assert labels == sorted(labels, key=rank.get)
    manifest = json.loads((folder/'data-manifest.json').read_text())
    assert any('ordem de leitura' in w for w in manifest['warnings'])


def prose(n, label):
    return ' '.join(
        f'This is independent document {label} about a distinct topic number {i} with original facts.'
        for i in range(n)
    )


def test_max_chars_filters_examples_not_whole_books(client):
    from app.studio.storage import save_source
    sources = [save_source(f'book{i}.txt', prose(40, i).encode())['id'] for i in range(4)]
    job = client.post('/api/studio/prepare', json={
        'source_ids': sources, 'min_chars': 1, 'max_chars': 180, 'min_words': 8, 'max_words': 20,
        'chunking': 'sentence', 'quality_filter': 'off'
    }).json()
    prepare(job, store.run_dir(job['id']))
    rows = json.loads((store.run_dir(job['id'])/'preview.json').read_text())
    assert len(rows) > 4
    assert all(len(row['text']) <= 180 for row in rows)
    assert len({row['group'] for row in rows}) >= 3


def test_drop_boilerplate_keeps_prose_in_same_split(client):
    from app.studio.storage import save_source
    catalog = 'Texto-fonte: Obra Completa, Editora Nova Aguilar. Publicado originalmente em folhetins na Revista Brasileira.'
    sources = [save_source('catalog.txt', catalog.encode())['id']]
    sources += [save_source(f'prose{i}.txt', prose(8, i).encode())['id'] for i in range(4)]
    job = client.post('/api/studio/prepare', json={
        'source_ids': sources, 'min_chars': 1, 'min_words': 8, 'quality_filter': 'off', 'drop_boilerplate': True
    }).json()
    prepare(job, store.run_dir(job['id']))
    rows = json.loads((store.run_dir(job['id'])/'preview.json').read_text())
    assert not any('Texto-fonte' in row['text'] for row in rows)
    assert any('independent document' in row['text'] for row in rows)
    related = [row for row in rows if row['metadata']['path'] == 'prose0.txt']
    assert related and len({row['split'] for row in related}) == 1


def test_llm_classifier_updates_metadata_without_storing_key(client, monkeypatch):
    from urllib.parse import urlsplit
    from app.studio.storage import save_source
    monkeypatch.delenv('OPENROUTER_API_KEY', raising=False)
    monkeypatch.setenv('LLM_API_KEY', 'secret-test-key')
    monkeypatch.setattr('app.studio.llm.public_target', lambda url: (urlsplit(url), '1.1.1.1'))
    monkeypatch.setattr('app.studio.network.public_target', lambda url: (urlsplit(url), '1.1.1.1'))

    def fake_complete(url, payload, timeout=45):
        items = json.loads(payload['messages'][1]['content'])['items']
        return json.dumps({'items': [
            {'id': item['id'], 'quality': 5, 'role': 'body', 'language': 'en', 'topic': 'training data'}
            for item in items
        ]})

    monkeypatch.setattr('app.studio.llm.chat_completions', fake_complete)
    sources = [save_source(f'doc{i}.txt', prose(8, i).encode())['id'] for i in range(4)]
    job = client.post('/api/studio/prepare', json={
        'source_ids': sources, 'min_chars': 1, 'min_words': 8, 'quality_filter': 'off', 'llm_classify': True
    }).json()
    folder = store.run_dir(job['id'])
    prepare(job, folder)
    dumped = json.dumps(json.loads((folder/'data-manifest.json').read_text()))
    rows = json.loads((folder/'preview.json').read_text())
    assert 'secret-test-key' not in dumped
    assert 'secret-test-key' not in json.dumps(rows)
    assert all(row['metadata']['topic'] == 'training data' for row in rows)
    assert all(row['metadata']['quality'] == 1.0 for row in rows)
    store.update('jobs', job['id'], {'status': 'completed'})
    preview = client.get('/api/studio/preparations/'+job['id']+'?role=body')
    assert preview.status_code == 200 and preview.json()['total'] == len(rows)


def test_llm_classify_requires_key(client, monkeypatch):
    from app.studio.storage import save_source
    monkeypatch.delenv('LLM_API_KEY', raising=False)
    monkeypatch.delenv('OPENROUTER_API_KEY', raising=False)
    sources = [save_source(f'doc{i}.txt', prose(8, i).encode())['id'] for i in range(4)]
    job = client.post('/api/studio/prepare', json={
        'source_ids': sources, 'min_chars': 1, 'llm_classify': True
    }).json()
    with pytest.raises(ValueError, match='LLM_API_KEY'):
        prepare(job, store.run_dir(job['id']))

