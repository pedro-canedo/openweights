"""Format adapters. Imported code is always data, never executed."""
import csv
import fnmatch
import io
import json
import re
import stat
import subprocess
import tempfile
import tokenize
import zipfile
from pathlib import Path, PurePosixPath

from .. import data

CODE = {'.py', '.js', '.ts', '.jsx', '.tsx', '.java', '.c', '.h', '.cpp', '.hpp', '.cs', '.go', '.rs', '.rb', '.php', '.swift', '.kt', '.scala', '.sql', '.sh', '.ps1', '.r', '.lua', '.vue', '.svelte', '.css', '.scss', '.ipynb'}
TEXT = {'.txt', '.md', '.markdown', '.rst', '.html', '.htm', '.xml', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.json', '.jsonl', '.csv'}
FORMATS = CODE | TEXT | {'.pdf', '.docx', '.zip'}
IGNORE = {'.git', 'node_modules', 'venv', '.venv', 'dist', 'build', 'target', '__pycache__', '.idea', '.next', '.cache'}
CONFIG = {'.yaml', '.yml', '.toml', '.ini', '.cfg', '.xml', '.json'}


def safe_name(name):
    name = name.replace('\\', '/')
    p = PurePosixPath(name)
    if p.is_absolute() or '..' in p.parts or ':' in name or any(ord(c) < 32 for c in name) or len(name) > 512 or not p.name:
        raise ValueError('Nome de arquivo ou caminho inseguro.')
    return str(p)


def ignored(name, recipe):
    p = PurePosixPath(name)
    if any(part in IGNORE or part == '.env' or part.startswith('.env.') for part in p.parts):
        return True
    if len(p.parts) > recipe.max_depth or any(fnmatch.fnmatch(name, pattern) for pattern in recipe.exclude):
        return True
    suffix = p.suffix.lower()
    if recipe.extensions and suffix not in recipe.extensions:
        return True
    if not recipe.include_docs and (suffix in {'.md', '.rst', '.txt'} or 'docs' in p.parts):
        return True
    if not recipe.include_tests and (any(x in {'test', 'tests', '__tests__'} for x in p.parts) or p.name.startswith('test_') or '.test.' in p.name or '.spec.' in p.name):
        return True
    return not recipe.include_config and suffix in CONFIG


def zip_members(raw, recipe):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > recipe.max_files:
            raise ValueError('ZIP excede a quantidade máxima de arquivos.')
        if sum(e.file_size for e in entries) > recipe.expanded_mb * 1024**2:
            raise ValueError('ZIP expandido excede o limite de tamanho.')
        for entry in entries:
            safe_name(entry.filename)
            if stat.S_ISLNK(entry.external_attr >> 16) or entry.flag_bits & 1:
                raise ValueError('ZIP com link simbólico ou criptografia não é aceito.')
            if entry.file_size > 32 * 1024**2 or entry.file_size / max(1, entry.compress_size) > 200:
                raise ValueError('ZIP contém arquivo excessivo ou taxa de compressão insegura.')
        for entry in entries:
            if not entry.is_dir():
                if ignored(entry.filename, recipe):
                    yield entry.filename, None
                else:
                    yield entry.filename, archive.read(entry)


def decode(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16'), 'utf-16'
    if b'\x00' in raw or raw.startswith((b'MZ', b'\x7fELF')):
        raise ValueError('Arquivo binário ou executável ignorado.')
    try:
        return raw.decode('utf-8-sig'), 'utf-8'
    except UnicodeDecodeError:
        return raw.decode('cp1252'), 'cp1252'


def extract(raw, name, recipe):
    suffix = Path(name).suffix.lower()
    if Path(name).name.lower() in {'readme', 'license', 'licence', 'copying', 'notice', 'dockerfile', 'makefile', '.gitignore', '.dockerignore'}:
        suffix = '.txt'
    if suffix not in FORMATS or suffix == '.zip':
        raise ValueError('Formato não suportado ou arquivo compactado aninhado.')
    if suffix == '.pdf':
        if not raw.startswith(b'%PDF-'):
            raise ValueError('Assinatura PDF inválida.')
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        total = len(reader.pages)
        cap = recipe.max_pdf_pages
        warning = None
        if total > cap:
            warning = f'PDF com {total} páginas; extraídas as primeiras {cap} (limite da receita). O restante não entra nesta preparação.'
        pages = []
        for number, page in enumerate(reader.pages[:cap], 1):
            text = page.extract_text() or ''
            if not text.strip() and recipe.ocr:
                with tempfile.TemporaryDirectory(prefix='ows-ocr-') as temporary:
                    source = Path(temporary) / 'input.pdf'
                    source.write_bytes(raw)
                    image = Path(temporary) / 'page'
                    subprocess.run(['pdftoppm', '-f', str(number), '-l', str(number), '-scale-to', '2400', '-singlefile', '-png', str(source), str(image)], check=True, timeout=60, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                    text = subprocess.run(['tesseract', str(image)+'.png', 'stdout', '-l', recipe.ocr_language], check=True, timeout=60, capture_output=True, text=True).stdout
            if text.strip():
                pages.append({'text': text, 'page': number})
        if not pages:
            raise ValueError('PDF sem texto. Ative OCR para páginas digitalizadas.')
        joined = '\n\n'.join(page['text'] for page in pages)
        row = {'text': joined, 'pages': [page['page'] for page in pages], 'pdf_pages': total}
        if warning:
            row['warning'] = warning
        return [row], 'pdf', 'binary'
    if suffix == '.docx':
        # Check expansion before python-docx opens the OOXML archive.
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 5000 or sum(e.file_size for e in archive.infolist()) > 128 * 1024**2:
                raise ValueError('DOCX expandido excede os limites.')
        from docx import Document
        doc = Document(io.BytesIO(raw))
        text = '\n\n'.join(p.text for p in doc.paragraphs)
        if recipe.tables:
            text += '\n\n' + '\n'.join(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return [{'text': text}], 'document', 'binary'
    text, encoding = decode(raw)
    if suffix in {'.html', '.htm'} and recipe.clean_html:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, 'html.parser')
        for element in soup(['script', 'style', 'noscript', 'template'] + (['nav', 'header', 'footer', 'aside'] if recipe.boilerplate else [])):
            element.decompose()
        text = soup.get_text('\n', strip=True)
    if suffix == '.xml':
        from defusedxml.ElementTree import fromstring
        text = '\n'.join(fromstring(text).itertext())
    if suffix in {'.json', '.jsonl', '.csv'}:
        if suffix == '.json':
            obj = json.loads(text)
            items = obj if isinstance(obj, list) else obj.get('data', obj.get('examples', [obj])) if isinstance(obj, dict) else [obj]
        elif suffix == '.jsonl':
            items = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            items = list(csv.DictReader(io.StringIO(text)))
        rows = []
        for obj in items:
            if isinstance(obj, dict) and recipe.instruction_field in obj and recipe.output_field in obj:
                obj = dict(obj, instruction=obj[recipe.instruction_field], output=obj[recipe.output_field], input=obj.get(recipe.input_field, ''))
            try:
                rows.append(data.normalize(obj))
            except ValueError:
                if recipe.preset == 'conversations':
                    raise
                rows.append({'text': json.dumps(obj, ensure_ascii=False)})
        return rows, 'structured', encoding
    if suffix == '.py' and not recipe.include_comments:
        tokens = tokenize.generate_tokens(io.StringIO(text).readline)
        text = tokenize.untokenize(t for t in tokens if t.type != tokenize.COMMENT)
    elif suffix in CODE and not recipe.include_comments:
        raise ValueError('Remoção segura de comentários disponível somente para Python. Ative incluir comentários para esta linguagem.')
    return [{'text': text}], 'code' if suffix in CODE else 'document', encoding
